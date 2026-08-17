# -*- coding: utf-8 -*-
import base64
import csv
import io

from odoo import Command, api, fields, models, _
from odoo.exceptions import UserError, ValidationError
from odoo.tools.float_utils import float_compare, float_is_zero

# TK công nợ: mỗi đối tác một dòng; bắt buộc có partner.
PARTNER_REQUIRED_PREFIXES = ('131', '331', '141', '138', '338')


class VasOpeningBalance(models.Model):
    """Bộ số dư đầu kỳ — nháp dùng chung cho import Excel và nhập tay."""

    _name = 'vas.opening.balance'
    _description = 'Số dư đầu kỳ VAS'
    _order = 'id desc'

    name = fields.Char(string='Số phiếu', required=True, copy=False, default='/')
    company_id = fields.Many2one(
        'res.company', required=True, index=True,
        default=lambda self: self.env.company, ondelete='cascade',
    )
    regime_id = fields.Many2one(
        'vas.regime', string='Chế độ kế toán', required=True,
        default=lambda self: self.env.company.vas_regime_id,
    )
    date = fields.Date(
        string='Ngày mở sổ',
        required=True,
        default=lambda self: self.env.company.vas_start_date,
        help='Phải trùng Ngày bắt đầu ghi sổ VAS của công ty.',
    )
    state = fields.Selection([
        ('draft', 'Nháp'),
        ('posted', 'Đã xác nhận'),
    ], default='draft', required=True, copy=False, index=True)
    ref = fields.Char(string='Diễn giải', default='Số dư đầu kỳ')
    line_ids = fields.One2many('vas.opening.balance.line', 'opening_id', string='Dòng số dư', copy=True)
    move_id = fields.Many2one('vas.move', string='Bút toán opening', copy=False, readonly=True)
    currency_id = fields.Many2one(
        'res.currency', default=lambda self: self.env.ref('base.VND'), required=True,
    )
    amount_debit = fields.Monetary(
        string='Tổng Nợ', compute='_compute_totals', currency_field='currency_id',
    )
    amount_credit = fields.Monetary(
        string='Tổng Có', compute='_compute_totals', currency_field='currency_id',
    )
    amount_diff = fields.Monetary(
        string='Lệch', compute='_compute_totals', currency_field='currency_id',
    )

    @api.depends('line_ids.debit', 'line_ids.credit')
    def _compute_totals(self):
        for rec in self:
            debit = sum(rec.line_ids.mapped('debit'))
            credit = sum(rec.line_ids.mapped('credit'))
            rec.amount_debit = debit
            rec.amount_credit = credit
            rec.amount_diff = debit - credit

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            company = self.env['res.company'].browse(
                vals.get('company_id') or self.env.company.id
            )
            if not vals.get('date') and company.vas_start_date:
                vals['date'] = company.vas_start_date
            if not vals.get('regime_id') and company.vas_regime_id:
                vals['regime_id'] = company.vas_regime_id.id
        return super().create(vals_list)

    @api.model
    def _partner_required_for_account(self, account):
        code = (account.code or '')
        return any(code.startswith(p) for p in PARTNER_REQUIRED_PREFIXES)

    def _validate_draft_lines(self):
        self.ensure_one()
        errors = []
        for idx, line in enumerate(self.line_ids, start=1):
            if not line.account_id:
                errors.append(_('Dòng %(n)s: thiếu tài khoản.', n=idx))
                continue
            if self._partner_required_for_account(line.account_id) and not line.partner_id:
                errors.append(_(
                    'Dòng %(n)s: TK %(code)s là công nợ — bắt buộc chọn đối tác.',
                    n=idx, code=line.account_id.code,
                ))
            if float_is_zero(line.debit, 2) and float_is_zero(line.credit, 2):
                errors.append(_('Dòng %(n)s: Nợ và Có đều = 0.', n=idx))
            if line.debit and line.credit:
                errors.append(_('Dòng %(n)s: không được vừa Nợ vừa Có.', n=idx))
        if errors:
            raise UserError('\n'.join(errors))

    def action_confirm(self):
        """Xác nhận nháp → POST vas.move opening tại vas_start_date."""
        for rec in self:
            if rec.state != 'draft':
                raise UserError(_('Chỉ bộ nháp mới xác nhận được.'))
            company = rec.company_id
            if not company.vas_start_date:
                raise UserError(_(
                    'Công ty chưa khai Ngày bắt đầu ghi sổ VAS — không xác nhận số dư đầu kỳ.'
                ))
            if rec.date != company.vas_start_date:
                raise UserError(_(
                    'Ngày mở sổ (%(d)s) phải trùng Ngày bắt đầu ghi sổ (%(c)s).',
                    d=rec.date, c=company.vas_start_date,
                ))
            if not rec.line_ids:
                raise UserError(_('Chưa có dòng số dư.'))
            rec._validate_draft_lines()
            if float_compare(rec.amount_debit, rec.amount_credit, precision_digits=2) != 0:
                raise UserError(_(
                    'Tổng Nợ (%(debit)s) khác Tổng Có (%(credit)s). Lệch = %(diff)s. Không post.',
                    debit=rec.amount_debit,
                    credit=rec.amount_credit,
                    diff=rec.amount_diff,
                ))
            # Cutoff + miễn trừ source_model nằm trong move.action_post.
            move = rec._create_opening_move()
            if rec.name == '/':
                rec.name = f'SDK/{rec.date}/{rec.id}'
            rec.write({'state': 'posted', 'move_id': move.id})
        return True

    def _create_opening_move(self):
        self.ensure_one()
        journal = self.env['vas.journal'].search([
            ('company_id', '=', self.company_id.id),
            ('code', '=', 'MOSO'),
            ('regime_id', '=', self.regime_id.id),
        ], limit=1) or self.env['vas.journal'].search([
            ('company_id', '=', self.company_id.id),
            ('code', '=', 'MOSO'),
        ], limit=1)
        if not journal:
            raise UserError(_('Thiếu sổ nhật ký VAS MOSO (mở sổ).'))
        lines = []
        seq = 10
        for line in self.line_ids.sorted(lambda l: (l.account_id.code or '', l.id)):
            lines.append(Command.create({
                'sequence': seq,
                'account_id': line.account_id.id,
                'name': self.ref or _('Số dư đầu kỳ'),
                'debit': line.debit,
                'credit': line.credit,
                'partner_id': line.partner_id.id if line.partner_id else False,
                'currency_id': self.currency_id.id,
            }))
            seq += 10
        move = self.env['vas.move'].create({
            'date': self.date,
            'journal_id': journal.id,
            'regime_id': self.regime_id.id,
            'move_kind': 'opening',
            'ref': self.ref or _('Số dư đầu kỳ'),
            'source_model': self._name,
            'source_res_id': self.id,
            'source_ref': self.display_name,
            'company_id': self.company_id.id,
            'currency_id': self.currency_id.id,
            'line_ids': lines,
        })
        move.action_post()
        return move

    def action_reopen(self):
        """Đảo opening đã post → mở lại nháp (chỉ khi kỳ còn mở; missing không cho §8.5)."""
        for rec in self:
            if rec.state != 'posted' or not rec.move_id:
                raise UserError(_('Chỉ bộ đã xác nhận mới mở lại được.'))
            # Kiểm kỳ khóa trên ngày gốc; post đảo đi qua action_post + miễn trừ.
            rec.env['vas.period']._assert_date_writable(
                rec.company_id,
                rec.move_id.date,
                doc_name=rec.move_id.name,
                allow_missing=False,
                move=rec.move_id,
            )
            rec.move_id.action_reverse()
            rec.write({'state': 'draft', 'move_id': False})
        return True

    def action_open_import_wizard(self):
        self.ensure_one()
        if self.state != 'draft':
            raise UserError(_('Chỉ import vào bộ nháp.'))
        return {
            'type': 'ir.actions.act_window',
            'name': _('Import số dư đầu kỳ'),
            'res_model': 'vas.opening.balance.import.wizard',
            'view_mode': 'form',
            'target': 'new',
            'context': {'default_opening_id': self.id},
        }


class VasOpeningBalanceLine(models.Model):
    _name = 'vas.opening.balance.line'
    _description = 'Dòng số dư đầu kỳ VAS'
    _order = 'account_id, partner_id, id'

    opening_id = fields.Many2one(
        'vas.opening.balance', required=True, ondelete='cascade', index=True,
    )
    account_id = fields.Many2one(
        'vas.account', string='Tài khoản', required=True, ondelete='restrict',
        domain="[('regime_id', '=', regime_id), ('active', '=', True)]",
    )
    partner_id = fields.Many2one('res.partner', string='Đối tác', ondelete='restrict')
    debit = fields.Monetary(string='Nợ', currency_field='currency_id', default=0.0)
    credit = fields.Monetary(string='Có', currency_field='currency_id', default=0.0)
    currency_id = fields.Many2one(related='opening_id.currency_id')
    company_id = fields.Many2one(related='opening_id.company_id', store=True)
    regime_id = fields.Many2one(related='opening_id.regime_id')

    @api.constrains('account_id', 'partner_id', 'debit', 'credit')
    def _check_line(self):
        for line in self:
            if line.opening_id.state != 'draft':
                raise ValidationError(_('Không sửa dòng khi bộ số dư đã xác nhận.'))
            if line.account_id and line.opening_id._partner_required_for_account(line.account_id):
                if not line.partner_id:
                    raise ValidationError(_(
                        'TK %(code)s là công nợ — bắt buộc chọn đối tác.',
                        code=line.account_id.code,
                    ))


class VasOpeningBalanceImportWizard(models.TransientModel):
    _name = 'vas.opening.balance.import.wizard'
    _description = 'Import Excel/CSV số dư đầu kỳ'

    opening_id = fields.Many2one('vas.opening.balance', required=True, ondelete='cascade')
    data_file = fields.Binary(string='File CSV/Excel', required=True)
    filename = fields.Char()
    mode = fields.Selection([
        ('replace', 'Thay toàn bộ nháp'),
        ('append', 'Thêm vào nháp hiện có'),
    ], default='replace', required=True)
    template_file = fields.Binary(string='Mẫu trống', compute='_compute_template')
    template_filename = fields.Char(default='mau_so_du_dau_ky.csv')

    @api.depends()
    def _compute_template(self):
        content = 'Mã TK,Đối tác,Nợ,Có\n111,,0,0\n131,Ten doi tac,0,0\n'
        data = base64.b64encode(content.encode('utf-8-sig'))
        for wiz in self:
            wiz.template_file = data

    def action_download_template(self):
        self.ensure_one()
        return {
            'type': 'ir.actions.act_url',
            'url': (
                f'/web/content/?model={self._name}&id={self.id}'
                f'&field=template_file&filename_field=template_filename&download=true'
            ),
            'target': 'self',
        }

    def _parse_rows(self):
        self.ensure_one()
        raw = base64.b64decode(self.data_file)
        name = (self.filename or '').lower()
        if name.endswith('.xlsx'):
            return self._parse_xlsx(raw)
        # CSV / Excel-saved CSV (default)
        text = raw.decode('utf-8-sig')
        reader = csv.DictReader(io.StringIO(text))
        if not reader.fieldnames:
            raise UserError(_('File trống hoặc thiếu header.'))
        # Normalize headers
        mapping = {}
        for h in reader.fieldnames:
            key = (h or '').strip().lower()
            if key in ('mã tk', 'ma tk', 'tk', 'code', 'account'):
                mapping[h] = 'code'
            elif key in ('đối tác', 'doi tac', 'partner', 'partner_id'):
                mapping[h] = 'partner'
            elif key in ('nợ', 'no', 'debit', 'dr'):
                mapping[h] = 'debit'
            elif key in ('có', 'co', 'credit', 'cr'):
                mapping[h] = 'credit'
        need = {'code', 'debit', 'credit'}
        if need - set(mapping.values()):
            raise UserError(_(
                'Header phải có cột: Mã TK, Đối tác, Nợ, Có. Nhận được: %(h)s',
                h=', '.join(reader.fieldnames),
            ))
        rows = []
        for i, row in enumerate(reader, start=2):
            parsed = {mapping[h]: (row.get(h) or '').strip() for h in mapping}
            if not any(parsed.values()):
                continue
            rows.append((i, parsed))
        return rows

    def _parse_xlsx(self, raw):
        try:
            from openpyxl import load_workbook
        except ImportError as err:
            raise UserError(_(
                'Không đọc được .xlsx (thiếu openpyxl). Hãy lưu file dạng CSV UTF-8.'
            )) from err
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if not header:
            raise UserError(_('File Excel trống.'))
        # Build fake CSV dicts
        headers = [str(h or '').strip() for h in header]
        fake = io.StringIO()
        writer = csv.writer(fake)
        writer.writerow(headers)
        for row in rows_iter:
            writer.writerow([('' if c is None else c) for c in row])
        fake.seek(0)
        # Reuse CSV parser via temporary binary
        self.data_file = base64.b64encode(fake.getvalue().encode('utf-8-sig'))
        self.filename = (self.filename or 'import') + '.csv'
        return self._parse_rows()

    def _find_account(self, code):
        code = (code or '').strip()
        return self.env['vas.account'].search([
            ('regime_id', '=', self.opening_id.regime_id.id),
            ('code', '=', code),
            ('active', '=', True),
        ], limit=1)

    def _find_partner(self, name):
        name = (name or '').strip()
        if not name:
            return self.env['res.partner']
        Partner = self.env['res.partner']
        partner = Partner.search([('name', '=', name)], limit=1)
        if not partner:
            partner = Partner.search([('ref', '=', name)], limit=1)
        if not partner:
            partner = Partner.search([('vat', '=', name)], limit=1)
        return partner

    def _to_float(self, value):
        if value in (None, ''):
            return 0.0
        if isinstance(value, (int, float)):
            return float(value)
        text = str(value).strip().replace(' ', '').replace(',', '')
        if not text:
            return 0.0
        return float(text)

    def action_import(self):
        self.ensure_one()
        opening = self.opening_id
        if opening.state != 'draft':
            raise UserError(_('Chỉ import vào bộ nháp.'))
        rows = self._parse_rows()
        if not rows:
            raise UserError(_('Không có dòng dữ liệu để import.'))

        errors = []
        vals_list = []
        Opening = self.env['vas.opening.balance']
        for row_no, parsed in rows:
            code = parsed.get('code') or ''
            account = self._find_account(code)
            if not account:
                errors.append(_('Dòng %(n)s: mã TK %(code)s không có trong vas.account.',
                                n=row_no, code=code or '(trống)'))
                continue
            partner_name = parsed.get('partner') or ''
            partner = self._find_partner(partner_name)
            if Opening._partner_required_for_account(account):
                if not partner_name:
                    errors.append(_(
                        'Dòng %(n)s: TK %(code)s là công nợ nhưng thiếu Đối tác.',
                        n=row_no, code=account.code,
                    ))
                    continue
                if not partner:
                    errors.append(_(
                        'Dòng %(n)s: không tìm thấy đối tác %(p)s.',
                        n=row_no, p=partner_name,
                    ))
                    continue
            elif partner_name and not partner:
                errors.append(_(
                    'Dòng %(n)s: không tìm thấy đối tác %(p)s.',
                    n=row_no, p=partner_name,
                ))
                continue
            try:
                debit = self._to_float(parsed.get('debit'))
                credit = self._to_float(parsed.get('credit'))
            except ValueError:
                errors.append(_('Dòng %(n)s: Nợ/Có không phải số.', n=row_no))
                continue
            vals_list.append({
                'opening_id': opening.id,
                'account_id': account.id,
                'partner_id': partner.id if partner else False,
                'debit': debit,
                'credit': credit,
            })

        if errors:
            raise UserError(_(
                'Import bị từ chối — không nhập dòng lỗi, không đoán:\n%(err)s',
                err='\n'.join(errors),
            ))
        if self.mode == 'replace':
            opening.line_ids.unlink()
        self.env['vas.opening.balance.line'].create(vals_list)
        return {
            'type': 'ir.actions.act_window',
            'res_model': 'vas.opening.balance',
            'res_id': opening.id,
            'view_mode': 'form',
            'target': 'current',
        }
