# -*- coding: utf-8 -*-
"""Khung báo cáo VAS — lớp trình bày + hợp đồng dữ liệu (không đụng account_reports).

Hợp đồng get_report_data(options) → {
    meta: { title, form_code, company_name, period_label, currency_id, warning, requires_account? },
    columns: [{ name, label, type: 'string'|'monetary'|'date', align }],
    lines: [{
        id, label, level, values: [...aligned with columns],
        unfoldable, unfolded, parent_id, is_total, is_leaf?, is_section?, class?,
    }],
    options: echo of input options,
    checks: optional dict (F01 diffs, is_balanced, ...),
}
OWL chỉ đọc columns + lines (+ meta). Số liệu vẫn từ vas.books.mixin.
"""
import base64
import io
import json
from collections import defaultdict

from odoo import api, fields, models, _
from odoo.exceptions import UserError
from odoo.tools.float_utils import float_round
from odoo.tools.safe_eval import safe_eval

# Bốn mẫu BCTC đi qua cùng khung OWL + PDF/XLSX (không vá từng báo cáo).
FS_WIZARD_MODELS = {
    'B01a-DNN': 'vas.b01a.wizard',
    'B02-DNN': 'vas.b02.wizard',
    'B03-DNN': 'vas.b03.wizard',
    'B09-DNN': 'vas.b09.wizard',
}
FS_CLIENT_XMLIDS = {
    'B01a-DNN': 'connecta_vas.action_vas_b01a_client',
    'B02-DNN': 'connecta_vas.action_vas_b02_client',
    'B03-DNN': 'connecta_vas.action_vas_b03_client',
    'B09-DNN': 'connecta_vas.action_vas_b09_client',
}


class VasReportEngine(models.AbstractModel):
    _name = 'vas.report.engine'
    _description = 'VAS report presentation helpers'

    @api.model
    def get_filter_defaults(self):
        """Options mặc định cho thanh điều khiển OWL (kỳ gần nhất + danh sách TK)."""
        company = self.env.company
        today = fields.Date.context_today(self)
        Period = self.env['vas.period']
        period = Period.search([
            ('fiscalyear_id.company_id', '=', company.id),
            ('date_start', '<=', today),
            ('date_end', '>=', today),
        ], limit=1)
        if not period:
            period = Period.search([
                ('fiscalyear_id.company_id', '=', company.id),
            ], order='date_start desc', limit=1)
        periods = Period.search([
            ('fiscalyear_id.company_id', '=', company.id),
        ], order='date_start desc', limit=36)
        accounts = []
        if company.vas_regime_id:
            accounts = [{
                'id': a.id,
                'name': a.display_name,
                'code': a.code or '',
            } for a in self.env['vas.account'].search([
                ('regime_id', '=', company.vas_regime_id.id),
                ('active', '=', True),
            ], order='code', limit=800)]
        return {
            'company_id': company.id,
            'company_name': company.name or '',
            'period_from_id': period.id if period else False,
            'period_to_id': period.id if period else False,
            'hide_reversed': True,
            'account_id': False,
            'partner_id': False,
            'periods': [{
                'id': p.id,
                'name': p.name,
                'date_start': fields.Date.to_string(p.date_start) if p.date_start else False,
                'date_end': fields.Date.to_string(p.date_end) if p.date_end else False,
            } for p in periods],
            'accounts': accounts,
        }

    @api.model
    def build_account_hierarchy_lines(self, leaf_rows, money_keys, make_values):
        """Dựng cây cha–con từ các dòng LÁ (đã có số).

        :param leaf_rows: list dict có account_id/account_code/account_name + money_keys
        :param money_keys: tuple field tiền để rollup
        :param make_values: callable(row_dict) → list values theo columns
        :return: list line dict hợp đồng (chưa gồm dòng tổng toàn báo cáo)

        Ràng buộc: số trên cha = tổng LÁ hậu duệ; is_leaf đánh dấu nguồn cộng tổng báo cáo.
        """
        if not leaf_rows:
            return []

        Account = self.env['vas.account']
        leaf_by_id = {row['account_id']: row for row in leaf_rows}
        needed_ids = set(leaf_by_id)
        for acc in Account.browse(list(needed_ids)):
            parent = acc.parent_id
            while parent and parent.id not in needed_ids:
                needed_ids.add(parent.id)
                parent = parent.parent_id

        accounts = Account.browse(list(needed_ids))
        acc_map = {a.id: a for a in accounts}
        children = defaultdict(list)
        roots = []
        for acc in accounts.sorted(lambda a: (a.code or '', a.id)):
            if acc.parent_id and acc.parent_id.id in needed_ids:
                children[acc.parent_id.id].append(acc.id)
            else:
                roots.append(acc.id)

        zeros = {k: 0.0 for k in money_keys}
        rollup_cache = {}

        def rollup(aid):
            if aid in rollup_cache:
                return rollup_cache[aid]
            own = leaf_by_id.get(aid)
            kids = children.get(aid) or []
            amounts = dict(zeros)
            if own:
                for k in money_keys:
                    amounts[k] = float_round(amounts[k] + (own.get(k) or 0.0), 2)
            for kid in kids:
                child_amt = rollup(kid)
                for k in money_keys:
                    amounts[k] = float_round(amounts[k] + child_amt[k], 2)
            rollup_cache[aid] = amounts
            return amounts

        for rid in roots:
            rollup(rid)

        lines = []

        def emit(aid, level, parent_line_id):
            acc = acc_map[aid]
            kids = children.get(aid) or []
            is_leaf = not kids
            amounts = rollup(aid)
            if is_leaf:
                row = dict(leaf_by_id[aid])
            else:
                row = {
                    'account_id': aid,
                    'account_code': acc.code or '',
                    'account_name': acc.name or '',
                    **amounts,
                }
            line_id = f'acc_{aid}'
            lines.append({
                'id': line_id,
                'label': f"{row['account_code']} {row['account_name']}".strip(),
                'level': level,
                'values': make_values(row),
                'unfoldable': bool(kids),
                'unfolded': False,
                'parent_id': parent_line_id or False,
                'is_total': False,
                'is_leaf': is_leaf,
                'is_section': bool(kids),
                'class': 'o_vas_section' if kids and level == 0 else (
                    'o_vas_section' if kids else ''
                ),
            })
            for kid in sorted(kids, key=lambda i: (acc_map[i].code or '', i)):
                emit(kid, level + 1, line_id)

        for rid in sorted(roots, key=lambda i: (acc_map[i].code or '', i)):
            emit(rid, 0, False)
        return lines

    @api.model
    def render_xlsx_bytes(self, data):
        """Sinh XLSX từ đúng cây columns + lines của hợp đồng."""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Border, Side, Alignment, PatternFill
        except ImportError as err:
            raise UserError(_('Thiếu openpyxl — không xuất Excel được.')) from err

        meta = data.get('meta') or {}
        columns = data.get('columns') or []
        lines = data.get('lines') or []

        wb = Workbook()
        ws = wb.active
        ws.title = (meta.get('form_code') or 'VAS')[:31]
        thin = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin'),
        )
        bold = Font(bold=True)
        warn_fill = PatternFill('solid', fgColor='FFF3CD')

        ws['A1'] = meta.get('company_name') or ''
        ws['A2'] = _('Mẫu số: %s', meta.get('form_code') or '')
        ws['A3'] = meta.get('title') or ''
        ws['A3'].font = Font(bold=True, size=14)
        ws['A4'] = meta.get('period_label') or ''
        if meta.get('account_label'):
            ws['A5'] = meta['account_label']
            row_i = 6
        else:
            row_i = 5
        if meta.get('warning'):
            cell = ws.cell(row=row_i, column=1, value=meta['warning'])
            cell.font = Font(bold=True, color='9A6700')
            cell.fill = warn_fill
            row_i += 1
        row_i += 1

        for col, cdef in enumerate(columns, 1):
            cell = ws.cell(row=row_i, column=col, value=cdef.get('label') or cdef.get('name'))
            cell.font = bold
            cell.border = thin
            cell.alignment = Alignment(wrap_text=True, horizontal='center')
        row_i += 1

        for line in lines:
            values = line.get('values') or []
            for col, val in enumerate(values, 1):
                cell = ws.cell(row=row_i, column=col, value=val)
                cell.border = thin
                if line.get('is_total') or line.get('is_section') or line.get('unfoldable'):
                    cell.font = bold
                cdef = columns[col - 1] if col <= len(columns) else {}
                if cdef.get('align') == 'right':
                    cell.alignment = Alignment(horizontal='right')
            row_i += 1

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    @api.model
    def action_download_xlsx(self, data, filename):
        content = self.render_xlsx_bytes(data)
        attachment = self.env['ir.attachment'].create({
            'name': filename,
            'type': 'binary',
            'datas': base64.b64encode(content),
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })
        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'self',
        }

    @api.model
    def fs_form_title(self, form_code):
        """Tiêu đề in hoa trên tờ giấy — theo mã mẫu, không hardcode B01a."""
        return {
            'B01a-DNN': _('BẢNG CÂN ĐỐI KẾ TOÁN'),
            'B02-DNN': _('BÁO CÁO KẾT QUẢ HOẠT ĐỘNG KINH DOANH'),
            'B03-DNN': _('BÁO CÁO LƯU CHUYỂN TIỀN TỆ'),
            'B09-DNN': _('BẢN THUYẾT MINH BÁO CÁO TÀI CHÍNH'),
        }.get(form_code) or (form_code or _('BÁO CÁO VAS'))

    @api.model
    def fs_action_name(self, form_code):
        """Tên action / breadcrumb (không viết hoa toàn bộ)."""
        return {
            'B01a-DNN': _('Bảng cân đối kế toán'),
            'B02-DNN': _('Báo cáo kết quả kinh doanh'),
            'B03-DNN': _('Báo cáo lưu chuyển tiền tệ'),
            'B09-DNN': _('Bản thuyết minh BCTC'),
        }.get(form_code) or _('Báo cáo VAS')

    @api.model
    def action_open_fs_client(self, snap):
        """Mở bản đã lập bằng khung OWL dùng chung (không form Odoo)."""
        snap.ensure_one()
        xmlid = FS_CLIENT_XMLIDS.get(snap.form_code)
        if xmlid:
            action = dict(self.env['ir.actions.actions']._for_xml_id(xmlid))
        else:
            action = {
                'type': 'ir.actions.client',
                'tag': 'vas_report_client',
                'name': self.fs_action_name(snap.form_code),
            }
        raw_ctx = action.get('context') or {}
        if isinstance(raw_ctx, str):
            raw_ctx = safe_eval(raw_ctx, {'uid': self.env.uid}) or {}
        ctx = dict(raw_ctx)
        ctx.update({
            'report_model': FS_WIZARD_MODELS.get(
                snap.form_code, 'vas.b01a.wizard',
            ),
            'report_title': self.fs_action_name(snap.form_code),
            'snapshot_id': snap.id,
            'period_from_id': snap.period_from_id.id,
            'period_to_id': snap.period_to_id.id,
            'hide_reversed': snap.hide_reversed,
            'company_id': snap.company_id.id,
        })
        action['context'] = ctx
        action['name'] = self.fs_action_name(snap.form_code)
        # Bỏ id/xml_id để client không tải lại action DB và mất snapshot_id.
        action.pop('id', None)
        action.pop('xml_id', None)
        return action

    @api.model
    def action_export_xlsx_fs(self, options, form_code):
        """XLSX từ hợp đồng đã lưu (snapshot_id) hoặc lập rồi đọc."""
        wizard_model = FS_WIZARD_MODELS.get(form_code)
        if not wizard_model:
            raise UserError(_('Mẫu %s không xuất Excel qua khung VAS.') % form_code)
        data = self.env[wizard_model].get_report_data(options or {})
        meta = data.get('meta') or {}
        filename = '%s_%s.xlsx' % (
            (meta.get('form_code') or form_code).replace('-', '_'),
            meta.get('snapshot_id') or 'live',
        )
        return self.action_download_xlsx(data, filename)

    @api.model
    def action_export_pdf_from_data(self, data, filename=None):
        """PDF một tờ từ đúng cây meta/columns/lines — mẫu thứ bảy thừa hưởng."""
        meta = data.get('meta') or {}
        name = filename or '%s.pdf' % (
            (meta.get('form_code') or 'VAS').replace('-', '_')
        )
        wiz = self.env['vas.report.print.wizard'].create({
            'data_json': json.dumps(data, ensure_ascii=False, default=str),
            'filename': name,
        })
        return self.env.ref('connecta_vas.action_report_vas_fs').report_action(
            wiz, config=False,
        )

    @api.model
    def action_export_pdf_fs(self, options, form_code):
        wizard_model = FS_WIZARD_MODELS.get(form_code)
        if not wizard_model:
            raise UserError(_('Mẫu %s không in PDF qua khung VAS.') % form_code)
        data = self.env[wizard_model].get_report_data(options or {})
        return self.action_export_pdf_from_data(data)
